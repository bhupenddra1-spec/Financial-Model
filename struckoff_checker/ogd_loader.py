"""Load MCA "Company Master Data" files (data.gov.in open data) into a SQLite database.

    python -m struckoff_checker.ogd_loader FILE [FILE ...] [-o mca_master.sqlite3] [--replace]

Accepts the state-wise CSV / XLSX files published on data.gov.in (and any file whose headers
resemble them, e.g. CORPORATE_IDENTIFICATION_NUMBER, COMPANY_NAME, COMPANY_STATUS ...).
Files are streamed row by row, so the full all-India dataset (millions of rows) loads without
holding it in memory. Use the result with ``STRUCKOFF_PROVIDER=mca_db``.

The open data has no PAN, so it supports CIN and company-name search only, and it is only as
current as the files you load - the loader records the load date and the tool displays it.
"""
import argparse
import csv
import re
import sqlite3
import sys
from datetime import datetime

from .providers import FIELDS
from .validators import normalise, normalise_name

# normalised header (lower-case, non-alphanumerics -> "_") -> our field
HEADER_ALIASES = {
    "cin": ["corporate_identification_number", "cin", "cin_llpin", "llpin"],
    "company_name": ["company_name", "name", "company", "company_name_as_per_cin"],
    "status": ["company_status", "status", "company_status_for_efiling"],
    "company_class": ["company_class", "class_of_company"],
    "category": ["company_category", "category"],
    "sub_category": ["company_sub_category", "company_subcategory", "sub_category"],
    "date_of_incorporation": ["date_of_registration", "date_of_incorporation"],
    "roc_code": ["registrar_of_companies", "roc_code", "roc"],
    "paid_up_capital": ["paidup_capital", "paid_up_capital", "paid_up_capital_rs"],
    "registered_address": ["registered_office_address", "registered_address"],
    "email": ["email_addr", "email_id", "email", "email_address"],
    "last_agm": ["latest_year_annual_return", "date_of_last_agm", "last_agm"],
    "last_balance_sheet": ["latest_year_financial_statement", "date_of_balance_sheet", "last_balance_sheet"],
    "listed": ["whether_listed_or_not", "listed"],
    "pan": ["pan", "pan_number"],
}
BATCH = 5000


def _key(header):
    return re.sub(r"[^a-z0-9]+", "_", str(header or "").lower()).strip("_")


def _column_map(headers):
    """field -> column index for one file's header row."""
    keys = [_key(h) for h in headers]
    return {field: keys.index(alias) for field, aliases in HEADER_ALIASES.items()
            for alias in aliases if alias in keys}


def _clean(field, value):
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    text = str(value).strip()
    if text.lower() in ("nan", "none", "null", "na", "n/a"):
        return "" if field != "listed" else text
    if field == "date_of_incorporation" and re.fullmatch(r"\d{2}-\d{2}-\d{4}", text):
        return text.replace("-", "/")
    if field == "paid_up_capital" and re.fullmatch(r"\d+\.0+", text):
        return text.split(".")[0]
    return text


def _iter_rows(path):
    """Yield lists of cell values (header first), streaming."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            yield list(row)
        wb.close()
        return
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        yield from csv.reader(f)


def init_db(db, replace=False):
    if replace:
        db.execute("DROP TABLE IF EXISTS companies")
        db.execute("DROP TABLE IF EXISTS meta")
    cols = ", ".join(f"{f} TEXT" + (" PRIMARY KEY" if f == "cin" else "") for f in FIELDS)
    db.execute(f"CREATE TABLE IF NOT EXISTS companies ({cols}, name_key TEXT)")
    db.execute("CREATE TABLE IF NOT EXISTS meta (key TEXT PRIMARY KEY, value TEXT)")


def load_file(db, path):
    """Load one file. Returns (rows_loaded, rows_skipped_without_cin)."""
    rows = _iter_rows(path)
    header = next(rows, None)
    if header is None:
        return 0, 0
    cmap = _column_map(header)
    if "cin" not in cmap or "company_name" not in cmap:
        raise ValueError(f"{path}: no CIN / company-name column found (headers: {list(header)[:8]}...)")
    sql = (f"INSERT OR REPLACE INTO companies ({', '.join(FIELDS)}, name_key) "
           f"VALUES ({', '.join('?' * (len(FIELDS) + 1))})")
    loaded = skipped = 0
    batch = []
    for row in rows:
        rec = {f: _clean(f, row[i]) if i < len(row) else "" for f, i in cmap.items()}
        cin = normalise(rec.get("cin"))
        if not cin or not rec.get("company_name"):
            skipped += 1
            continue
        rec["cin"] = cin
        rec["registration_number"] = cin[-6:] if len(cin) == 21 else ""  # CIN ends with the RoC serial
        rec["pan_name"] = rec.get("pan_name", "")
        batch.append([rec.get(f, "") for f in FIELDS] + [normalise_name(rec["company_name"])])
        loaded += 1
        if len(batch) >= BATCH:
            db.executemany(sql, batch)
            batch.clear()
    if batch:
        db.executemany(sql, batch)
    db.commit()
    return loaded, skipped


def finish(db, files):
    db.execute("CREATE INDEX IF NOT EXISTS idx_name ON companies(name_key)")
    db.execute("CREATE INDEX IF NOT EXISTS idx_pan ON companies(pan)")
    total = db.execute("SELECT COUNT(*) FROM companies").fetchone()[0]
    prior = db.execute("SELECT value FROM meta WHERE key='files'").fetchone()
    names = (prior[0].split("|") if prior else []) + [f.replace("\\", "/").rsplit("/", 1)[-1] for f in files]
    for k, v in (("loaded_at", datetime.now().strftime("%d-%b-%Y")), ("rows", str(total)),
                 ("files", "|".join(dict.fromkeys(names)))):
        db.execute("REPLACE INTO meta VALUES (?, ?)", (k, v))
    db.commit()
    return total


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("files", nargs="+", help="CSV/XLSX company master files (e.g. state-wise from data.gov.in)")
    ap.add_argument("-o", "--output", default="mca_master.sqlite3", help="SQLite file to create/extend")
    ap.add_argument("--replace", action="store_true", help="start from an empty database")
    args = ap.parse_args(argv)
    db = sqlite3.connect(args.output)
    init_db(db, args.replace)
    for path in args.files:
        try:
            loaded, skipped = load_file(db, path)
        except (OSError, ValueError) as exc:
            print(f"ERROR {exc}", file=sys.stderr)
            return 1
        print(f"{path}: {loaded} companies loaded, {skipped} rows skipped (no CIN/name)")
    print(f"Database {args.output}: {finish(db, args.files)} companies in total")
    return 0


if __name__ == "__main__":
    sys.exit(main())
