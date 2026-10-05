"""Reading supplier lists and writing the report in the 'Struck-Off Companies' workbook format.

Report sheets (same as the reference format):
  Suppliers              - the input list (s.no, Company Name, GST, PAN)
  Supplier-Final Sheet   - MCA master data for every supplier (+ struck-off flag)
  Supplier- Double Status- suppliers whose matches carry conflicting statuses
  Sheet1                 - PAN name vs. name in master reconciliation
  Summary                - counts per flag
"""
import csv
import difflib
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .providers import FIELD_LABELS
from .validators import normalise_name

# Layout of the "Supplier-Final Sheet": input name, PAN, then CIN + master fields.
FINAL_HEADERS = ["Company Name", "PAN"] + [label for _, label in FIELD_LABELS] + ["Struck-Off Check", "Remarks"]
FLAG_FILLS = {
    "Struck-Off": "FFC7CE", "Under Strike-Off Process": "FFD8A8", "Double Status": "FFEB9C",
    "Other Status": "FFEB9C", "Active": "C6EFCE", "Not Found": "D9D9D9",
    "Invalid Input": "FFC7CE", "Error": "FFC7CE",
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")

_HEADER_RE = {
    "name": re.compile(r"company|compamy|name|supplier|vendor"),
    "gstin": re.compile(r"\bgst"),
    "pan": re.compile(r"\bpan\b"),
    "cin": re.compile(r"\bcin\b|llpin"),
}
_NOT_NAME_RE = re.compile(r"\bgst|\bpan\b|\bcin\b|llpin|s\.?\s?no")


def _rows_from_upload(filename, data):
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb["Suppliers"] if "Suppliers" in wb.sheetnames else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith((".csv", ".txt")):
        return list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    raise ValueError("Upload an .xlsx, .csv or .txt file")


def extract_suppliers(filename, data):
    """Read supplier rows -> [{name, gstin, pan, cin}] using the header row to find columns."""
    rows = [r for r in _rows_from_upload(filename, data) if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    heads = [str(c or "").strip().lower() for c in rows[0]]
    cols = {}
    for field in ("gstin", "pan", "cin", "name"):  # name last: "company" also appears in other headers
        for i, h in enumerate(heads):
            if i not in cols.values() and _HEADER_RE[field].search(h) \
                    and not (field == "name" and _NOT_NAME_RE.search(h)):
                cols[field] = i
                break
    if not cols:  # headerless: treat a single column as company names
        rows, cols = [[]] + rows, {"name": 0}
    out = []
    for r in rows[1:]:
        item = {f: ("" if i >= len(r) or r[i] is None else str(r[i]).strip()) for f, i in cols.items()}
        item = {f: item.get(f, "") for f in ("name", "gstin", "pan", "cin")}
        if any(item.values()):
            out.append(item)
    return out


def _style_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def _autosize(ws, cap=60):
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), cap)


def _final_row(result, rec):
    inp = result["input"]
    if rec is None:
        values = [""] * len(FIELD_LABELS)
    else:
        values = [rec.get(k, "") for k, _ in FIELD_LABELS]
    return [inp["name"], inp["pan"] or (rec or {}).get("pan", "")] + values + [result["flag"], result["remarks"]]


def _name_match(a, b):
    if not a or not b:
        return ""
    x, y = normalise_name(a), normalise_name(b)
    if x == y:
        return "Match"
    return "Similar" if difflib.SequenceMatcher(None, x, y).ratio() >= 0.85 else "Mismatch"


def build_report(results, summary=None):
    wb = Workbook()

    ws = wb.active
    ws.title = "Suppliers"
    ws.append(["s.no", "Company Name", "GST", "PAN"])
    for n, r in enumerate(results, start=1):
        ws.append([n, r["input"]["name"], r["input"]["gstin"], r["input"]["pan"]])
    _style_header(ws)
    _autosize(ws)

    final = wb.create_sheet("Supplier-Final Sheet")
    final.append(FINAL_HEADERS)
    flag_col = FINAL_HEADERS.index("Struck-Off Check") + 1
    for r in results:
        for rec in (r["records"] or [None]):
            final.append(_final_row(r, rec))
            fill = FLAG_FILLS.get((rec or {}).get("group") if r["flag"] == "Double Status" else r["flag"])
            if fill:
                final.cell(final.max_row, flag_col).fill = PatternFill("solid", fgColor=fill)
    _style_header(final)
    _autosize(final, cap=50)
    final.auto_filter.ref = final.dimensions

    double = wb.create_sheet("Supplier- Double Status")
    double.append(FINAL_HEADERS)
    for r in results:
        if r["flag"] == "Double Status":
            for rec in r["records"]:
                double.append(_final_row(r, rec))
    _style_header(double)
    _autosize(double, cap=50)

    s1 = wb.create_sheet("Sheet1")
    s1.append(["S.No.", "PAN", "PAN Name As Per Traces", "Name In Master", "Name Check"])
    seen = set()
    for r in results:
        for rec in r["records"]:
            pan = rec.get("pan") or r["input"]["pan"]
            if not pan or pan in seen:
                continue
            seen.add(pan)
            s1.append([len(seen), pan, rec.get("pan_name", ""), rec.get("company_name", ""),
                       _name_match(rec.get("pan_name", ""), rec.get("company_name", ""))])
    _style_header(s1)
    _autosize(s1)

    if summary:
        sm = wb.create_sheet("Summary")
        sm.append(["Metric", "Count"])
        for k, v in summary.items():
            sm.append([k, v])
        _style_header(sm)
        sm.column_dimensions["A"].width = 28

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    ws.append(["s.no", "Company Name", "GST", "PAN", "CIN (optional)"])
    ws.append([1, "Example Supplier Private Limited", "27AAACA0242K1Z5", "AAACA0242K", ""])
    ws.append([2, "Another Supplier Limited", "", "", "L12345MH2000PLC123456"])
    for col, w in zip("ABCDE", (6, 48, 22, 16, 26)):
        ws.column_dimensions[col].width = w
    _style_header(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
