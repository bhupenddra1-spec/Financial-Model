"""Data providers that resolve a CIN / PAN / company name to MCA master data.

MCA's "View Company Master Data" page is CAPTCHA-protected and has no free API,
so (as with the MSME tool) live lookups go through a licensed data provider.

* ``demo``  - deterministic synthetic data for trying the tool.
* ``local`` - your own MCA master-data file (CSV / XLSX), loaded into memory (small files).
* ``mca_db`` - SQLite database built from data.gov.in files by ``ogd_loader`` (millions of rows).
* ``http``  - any REST API, configured through environment variables.

Every provider implements ``search(kind, value)`` -> list of record dicts keyed
by ``FIELDS`` (empty list when nothing matches). ``kind`` is CIN, PAN or NAME.
"""
import csv
import difflib
import hashlib
import json
import os
from datetime import date, timedelta

from .validators import normalise, normalise_name

# key -> header used in the Excel report ("Supplier-Final Sheet" layout)
FIELD_LABELS = [
    ("cin", "CIN"),
    ("company_name", "Company Name As Per CIN"),
    ("status", "Company Status(for efiling)"),
    ("roc_code", "ROC Code"),
    ("registration_number", "Registration Number"),
    ("category", "Company Category"),
    ("sub_category", "Company SubCategory"),
    ("company_class", "Class of Company"),
    ("paid_up_capital", "Paid up Capital(Rs)"),
    ("members", "Number of Members(Applicable in case of company without Share Capital)"),
    ("date_of_incorporation", "Date of Incorporation"),
    ("registered_address", "Registered Address"),
    ("other_address", "Address other than R/o where all or any books of account and papers are maintained"),
    ("email", "Email Id"),
    ("listed", "Whether Listed or not"),
    ("suspended", "Suspended at stock exchange"),
    ("suspended_detail", "Suspended at stock exchange"),
    ("last_agm", "Date of last AGM"),
    ("last_balance_sheet", "Date of Balance Sheet"),
]
FIELDS = [k for k, _ in FIELD_LABELS] + ["pan", "pan_name"]


class ProviderError(Exception):
    """A lookup could not be completed (network, auth, bad response, missing file)."""


class BaseProvider:
    name = "base"

    def search(self, kind, value):
        raise NotImplementedError


def _blank(record):
    return {f: ("" if record.get(f) is None else str(record.get(f))) for f in FIELDS}


class DemoProvider(BaseProvider):
    name = "demo"
    _STATUSES = ["Active"] * 12 + ["Strike Off"] * 3 + ["Under process of striking off"] * 2 + \
                ["Amalgamated", "Dormant Company", "Active"]
    _ROC = [("MH", "RoC-Mumbai"), ("DL", "RoC-Delhi"), ("KA", "RoC-Bangalore"),
            ("GJ", "RoC-Ahmedabad"), ("TN", "RoC-Chennai"), ("UP", "RoC-Kanpur")]

    def _record(self, h, name=None, pan=None, cin=None, status=None):
        st, roc = self._ROC[h[1] % len(self._ROC)]
        inc = date(1990, 1, 1) + timedelta(days=int.from_bytes(h[2:4], "big") % 11000)
        serial = int.from_bytes(h[4:7], "big") % 10**6
        cin = cin or f"U{int.from_bytes(h[7:9], 'big') % 100000:05d}{st}{inc.year}PTC{serial:06d}"
        name = name or f"{['SHREE GANESH', 'APEX', 'SUNRISE', 'BHARAT', 'VERTEX', 'PIONEER'][h[9] % 6]} " \
                       f"{['INDUSTRIES', 'POLYMERS', 'TECH SOLUTIONS', 'LOGISTICS'][h[10] % 4]} PRIVATE LIMITED"
        status = status or self._STATUSES[h[11] % len(self._STATUSES)]
        struck = "strik" in status.lower()
        return _blank({
            "cin": cin, "company_name": name, "status": status, "roc_code": roc,
            "registration_number": serial, "category": "Company limited by Shares",
            "sub_category": "Non-govt company", "company_class": "Private",
            "paid_up_capital": (h[12] % 90 + 1) * 100000, "date_of_incorporation": inc.strftime("%d/%m/%Y"),
            "registered_address": f"{h[13] % 300 + 1}, Demo Industrial Estate, {roc[4:]}",
            "email": f"info{serial}@example.com", "listed": "Unlisted", "suspended": "NA",
            "last_agm": "" if struck else "30/09/2025", "last_balance_sheet": "" if struck else "31/03/2025",
            "pan": pan or "", "pan_name": name,
        })

    def search(self, kind, value):
        h = hashlib.sha256(f"{kind}:{value}".encode()).digest()
        if h[0] % 16 == 0:
            return []
        if kind == "CIN":
            return [self._record(h, cin=value)]
        if kind == "PAN":
            return [self._record(h, pan=value)]
        recs = [self._record(h, name=value.upper())]
        if h[14] % 8 == 0:  # some names resolve to two entries with different statuses
            h2 = hashlib.sha256(h).digest()
            recs.append(self._record(h2, name=value.upper(),
                                     status="Strike Off" if "strik" not in recs[0]["status"].lower() else "Active"))
        return recs


class LocalFileProvider(BaseProvider):
    """Looks up a master file with one row per company (CSV or XLSX)."""

    name = "local"
    _ALIASES = {
        "cin": ["cin", "cin/llpin", "llpin", "corporate identity number"],
        "company_name": ["company_name", "company name", "company name as per cin", "name"],
        "status": ["status", "company status", "company status(for efiling)", "company status (for efiling)"],
        "roc_code": ["roc_code", "roc code", "roc"],
        "registration_number": ["registration_number", "registration number"],
        "category": ["category", "company category"],
        "sub_category": ["sub_category", "company subcategory", "subcategory"],
        "company_class": ["company_class", "class of company", "class"],
        "paid_up_capital": ["paid_up_capital", "paid up capital(rs)", "paid up capital", "paidup capital"],
        "members": ["members", "number of members(applicable in case of company without share capital)"],
        "date_of_incorporation": ["date_of_incorporation", "date of incorporation"],
        "registered_address": ["registered_address", "registered address", "registered office address"],
        "other_address": ["other_address", "address other than r/o where all or any books of account and papers are maintained"],
        "email": ["email", "email id", "email address"],
        "listed": ["listed", "whether listed or not"],
        "suspended": ["suspended", "suspended at stock exchange"],
        "last_agm": ["last_agm", "date of last agm"],
        "last_balance_sheet": ["last_balance_sheet", "date of balance sheet"],
        "pan": ["pan", "pan number"],
        "pan_name": ["pan_name", "pan name as per traces", "name as per pan"],
    }

    def __init__(self, path):
        self.records = [self._map(r) for r in self._read_rows(path)]
        self.by_cin, self.by_pan, self.by_name = {}, {}, {}
        for rec in self.records:
            if rec["cin"]:
                self.by_cin.setdefault(normalise(rec["cin"]), []).append(rec)
            if rec["pan"]:
                self.by_pan.setdefault(normalise(rec["pan"]), []).append(rec)
            self.by_name.setdefault(normalise_name(rec["company_name"]), []).append(rec)

    @staticmethod
    def _read_rows(path):
        if not os.path.exists(path):
            raise ProviderError(f"Company master file not found: {path}")
        if path.lower().endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            rows = list(load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True))
            if not rows:
                return []
            heads = [str(h or "").strip() for h in rows[0]]
            return [dict(zip(heads, r)) for r in rows[1:]]
        with open(path, newline="", encoding="utf-8-sig") as f:
            return list(csv.DictReader(f))

    def _map(self, row):
        low = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
        rec = {}
        for field, aliases in self._ALIASES.items():
            value = next((low[a] for a in aliases if low.get(a) not in (None, "")), "")
            rec[field] = value.strftime("%d/%m/%Y") if hasattr(value, "strftime") else str(value).strip()
        return _blank(rec)

    def search(self, kind, value):
        if kind == "CIN":
            return list(self.by_cin.get(value, []))
        if kind == "PAN":
            return list(self.by_pan.get(value, []))
        key = normalise_name(value)
        if key in self.by_name:
            return list(self.by_name[key])
        close = difflib.get_close_matches(key, self.by_name, n=1, cutoff=0.93)
        return list(self.by_name[close[0]]) if close else []


class MasterDbProvider(BaseProvider):
    """Reads the SQLite database built by ``python -m struckoff_checker.ogd_loader``.

    Handles millions of rows. Open data has no PAN, so PAN search is reported as
    unsupported (None) unless the loaded files carried a PAN column.
    """

    name = "mca_db"

    def __init__(self, path):
        import sqlite3
        if not os.path.exists(path):
            raise ProviderError(f"MCA database not found: {path} (build it with python -m struckoff_checker.ogd_loader)")
        self._db = sqlite3.connect(f"file:{path}?mode=ro", uri=True, check_same_thread=False)
        self._db.row_factory = sqlite3.Row
        try:
            self._has_pan = self._db.execute("SELECT 1 FROM companies WHERE pan != '' LIMIT 1").fetchone() is not None
            meta = dict(self._db.execute("SELECT key, value FROM meta").fetchall())
        except sqlite3.DatabaseError as exc:
            raise ProviderError(f"{path} is not a valid MCA database: {exc}") from exc
        self.info = f"MCA open data loaded {meta.get('loaded_at', '?')} ({int(meta.get('rows', 0)):,} companies)"

    def search(self, kind, value):
        if kind == "PAN" and not self._has_pan:
            return None
        column, key = {"CIN": ("cin", value), "PAN": ("pan", value), "NAME": ("name_key", normalise_name(value))}[kind]
        rows = self._db.execute(f"SELECT {', '.join(FIELDS)} FROM companies WHERE {column} = ? LIMIT 20", (key,)).fetchall()
        return [_blank(dict(r)) for r in rows]


class HttpApiProvider(BaseProvider):
    """Generic REST provider.

    STRUCKOFF_API_URL_CIN / _PAN / _NAME   URL templates containing ``{id}`` (URL-encoded)
    STRUCKOFF_API_METHOD                   GET (default) or POST (sends {"query": value, "type": kind})
    STRUCKOFF_API_KEY, _KEY_HEADER         auth value / header name (default Authorization)
    STRUCKOFF_API_RESULTS_PATH             dotted path to the list of matches (blank: response is the record)
    STRUCKOFF_API_FIELD_MAP                JSON: our field -> dotted path within one match
    """

    name = "http"

    def __init__(self, env=os.environ):
        self.urls = {k: env.get(f"STRUCKOFF_API_URL_{k}", "") for k in ("CIN", "PAN", "NAME")}
        if not any(self.urls.values()):
            raise ProviderError("Set STRUCKOFF_API_URL_CIN / _PAN / _NAME for the http provider")
        self.method = env.get("STRUCKOFF_API_METHOD", "GET").upper()
        self.timeout = float(env.get("STRUCKOFF_API_TIMEOUT", "20"))
        self.headers = {"Accept": "application/json"}
        if env.get("STRUCKOFF_API_KEY"):
            self.headers[env.get("STRUCKOFF_API_KEY_HEADER", "Authorization")] = env["STRUCKOFF_API_KEY"]
        self.results_path = env.get("STRUCKOFF_API_RESULTS_PATH", "")
        self.field_map = json.loads(env.get("STRUCKOFF_API_FIELD_MAP", "{}")) or {f: f for f in FIELDS}

    @staticmethod
    def _dig(data, path):
        for part in filter(None, path.split(".")):
            if isinstance(data, list) and part.isdigit():
                data = data[int(part)] if int(part) < len(data) else None
            elif isinstance(data, dict):
                data = data.get(part)
            else:
                return None
        return data

    def search(self, kind, value):
        import requests
        from urllib.parse import quote

        template = self.urls.get(kind)
        if not template:
            return None if kind != "NAME" else []  # caller treats None as "unsupported"
        try:
            if self.method == "POST":
                resp = requests.post(template.replace("{id}", quote(value)), json={"query": value, "type": kind},
                                     headers=self.headers, timeout=self.timeout)
            else:
                resp = requests.get(template.replace("{id}", quote(value)), headers=self.headers, timeout=self.timeout)
        except requests.RequestException as exc:
            raise ProviderError(f"API request failed: {exc}") from exc
        if resp.status_code == 404:
            return []
        if resp.status_code >= 400:
            raise ProviderError(f"API returned HTTP {resp.status_code}")
        try:
            payload = resp.json()
        except ValueError as exc:
            raise ProviderError("API returned a non-JSON response") from exc
        items = self._dig(payload, self.results_path) if self.results_path else payload
        if items in (None, "", {}):
            return []
        items = items if isinstance(items, list) else [items]
        out = []
        for item in items:
            rec = _blank({f: self._dig(item, p) for f, p in self.field_map.items()})
            if rec["cin"] or rec["company_name"]:
                out.append(rec)
        return out


def build_provider(env=os.environ):
    kind = env.get("STRUCKOFF_PROVIDER", "demo").lower()
    if kind == "demo":
        return DemoProvider()
    if kind == "local":
        return LocalFileProvider(env.get("STRUCKOFF_MASTER_FILE", "sample_data/struckoff_master.csv"))
    if kind == "mca_db":
        return MasterDbProvider(env.get("STRUCKOFF_MCA_DB", "mca_master.sqlite3"))
    if kind == "http":
        return HttpApiProvider(env)
    raise ProviderError(f"Unknown STRUCKOFF_PROVIDER '{kind}' (use demo, local, mca_db or http)")
