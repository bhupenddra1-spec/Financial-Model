"""Reading bulk input files and writing the formatted Excel report."""
import csv
import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .validators import detect_id_type

REPORT_COLUMNS = [
    ("query", "Input"),
    ("id_type", "ID Type"),
    ("status", "MSME Registration Status"),
    ("udyam_number", "Udyam Registration Number"),
    ("pan", "PAN"),
    ("enterprise_name", "Name of Enterprise"),
    ("enterprise_type", "Type of Enterprise"),
    ("major_activity", "Major Activity of Enterprise"),
    ("organisation_type", "Type of Organisation"),
    ("date_of_incorporation", "Date of Incorporation"),
    ("date_of_udyam_registration", "Date of Udyam Registration"),
    ("state", "State"),
    ("district", "District"),
    ("msme_payment_rule", "45-day rule / Sec 43B(h) applies"),
    ("remarks", "Remarks"),
    ("source", "Source"),
    ("checked_at", "Checked At"),
]

STATUS_FILLS = {
    "Registered": "C6EFCE",
    "Not Registered": "FFEB9C",
    "Invalid Format": "FFC7CE",
    "Error": "FFC7CE",
}


def _rows_from_upload(filename, data):
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    raise ValueError("Upload an .xlsx, .csv or .txt file")


def extract_identifiers(filename, data):
    """Pull PAN / Udyam numbers from an uploaded sheet.

    Uses a column whose header mentions PAN, Udyam or ID; otherwise the column
    with the most valid identifiers.
    """
    rows = [r for r in _rows_from_upload(filename, data) if any(c not in (None, "") for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [list(r) + [None] * (width - len(r)) for r in rows]
    header = [str(c or "").strip().lower() for c in rows[0]]
    header_col = next((i for i, h in enumerate(header)
                       if any(k in h for k in ("pan", "udyam", "urn", "id number", "identifier"))), None)
    if header_col is not None:
        return [str(r[header_col]).strip() for r in rows[1:] if r[header_col] not in (None, "")]
    scores = [sum(1 for r in rows if detect_id_type(r[i])) for i in range(width)]
    col = max(range(width), key=lambda i: scores[i])
    body = rows[1:] if not detect_id_type(rows[0][col]) else rows
    return [str(r[col]).strip() for r in body if r[col] not in (None, "")]


def build_report(results, summary=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "MSME Verification"
    ws.append([label for _, label in REPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    status_idx = [k for k, _ in REPORT_COLUMNS].index("status") + 1
    for r in results:
        ws.append([r.get(k, "") for k, _ in REPORT_COLUMNS])
        fill = STATUS_FILLS.get(r.get("status"))
        if fill:
            ws.cell(ws.max_row, status_idx).fill = PatternFill("solid", fgColor=fill)
    for i, (key, label) in enumerate(REPORT_COLUMNS, start=1):
        longest = max([len(label)] + [len(str(r.get(key, ""))) for r in results])
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, longest + 2), 48)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if summary:
        s = wb.create_sheet("Summary")
        s.append(["Metric", "Count"])
        for cell in s[1]:
            cell.font = Font(bold=True)
        for k, v in summary.items():
            s.append([k, v])
        s.column_dimensions["A"].width = 24

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"
    ws.append(["Vendor Code", "Vendor Name", "PAN / Udyam Number"])
    ws.append(["V001", "Example Supplier 1", "AAAPL1234C"])
    ws.append(["V002", "Example Supplier 2", "UDYAM-MH-26-0012345"])
    for col, w in zip("ABC", (14, 30, 28)):
        ws.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
