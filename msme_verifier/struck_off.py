"""Struck-off company identification: supplier PAN / GSTIN / CIN -> MCA master data.

MCA's "View Company/LLP Master Data" has no free API (it sits behind a CAPTCHA),
so lookups go through a provider, as with the MSME tool:

* ``demo``  - deterministic synthetic data, for trying the tool.
* ``local`` - a company master file (CSV / XLSX) you maintain, with a PAN column.
* ``http``  - a licensed company-data REST API, configured through MCA_API_* variables.
"""
import difflib
import hashlib
import io
import json
import os
import re
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .providers import ProviderError, api_headers, dig, fetch_json, read_table
from .validators import PAN_RE, normalise, pan_holder_type

GSTIN_RE = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]$")
CIN_RE = re.compile(r"^[LU][0-9]{5}[A-Z]{2}[0-9]{4}[A-Z]{3}[0-9]{6}$")
LLPIN_RE = re.compile(r"^[A-Z]{3}-[0-9]{4}$")

# MCA master-data fields, in the column order of the report's output sheets.
COMPANY_COLUMNS = [
    ("cin", "CIN"),
    ("company_name", "Company Name As Per CIN"),
    ("status", "Company Status(for efiling)"),
    ("roc_code", "ROC Code"),
    ("registration_number", "Registration Number"),
    ("category", "Company Category"),
    ("subcategory", "Company SubCategory"),
    ("class_of_company", "Class of Company"),
    ("paid_up_capital", "Paid up Capital(Rs)"),
    ("number_of_members", "Number of Members(Applicable in case of company without Share Capital)"),
    ("date_of_incorporation", "Date of Incorporation"),
    ("registered_address", "Registered Address"),
    ("other_address", "Address other than R/o where all or any books of account and papers are maintained"),
    ("email", "Email Id"),
    ("listed", "Whether Listed or not"),
    ("suspended", "Suspended at stock exchange"),
    ("active_compliance", "ACTIVE Compliance"),
    ("date_of_last_agm", "Date of last AGM"),
    ("date_of_balance_sheet", "Date of Balance Sheet"),
]
COMPANY_FIELDS = [k for k, _ in COMPANY_COLUMNS]

OUTCOME_ACTIVE = "Active"
OUTCOME_STRUCK_OFF = "Struck Off"
OUTCOME_UNDER_PROCESS = "Under Process of Striking Off"
OUTCOME_NOT_ACTIVE = "Other - Not Active"
OUTCOME_NOT_FOUND = "Not Found on MCA"
OUTCOME_NOT_COMPANY = "Not a Company / LLP"
OUTCOME_INVALID = "Invalid Input"
OUTCOME_ERROR = "Error"

RISK_FILLS = {"high": "FFC7CE", "medium": "FFEB9C", "ok": "C6EFCE"}


def classify_status(status):
    """Map an MCA 'Company Status (for efiling)' value to (outcome, risk)."""
    s = (status or "").strip().lower()
    if not s:
        return OUTCOME_NOT_FOUND, "na"
    if "under process" in s and "strik" in s:
        return OUTCOME_UNDER_PROCESS, "high"
    if "strike" in s or "struck" in s:
        return OUTCOME_STRUCK_OFF, "high"
    if s == "active" or s.startswith("active"):
        return OUTCOME_ACTIVE, "ok"
    if any(w in s for w in ("dissolved", "liquidated", "vanished")):
        return OUTCOME_NOT_ACTIVE, "high"
    return OUTCOME_NOT_ACTIVE, "medium"  # dormant, amalgamated, under liquidation, converted...


def name_similarity(a, b):
    def clean(n):
        n = (n or "").upper().replace(".", " ")
        n = re.sub(r"\bPVT\b", "PRIVATE", n)
        n = re.sub(r"\bLTD\b", "LIMITED", n)
        n = re.sub(r"[^A-Z0-9 ]", " ", n)
        return " ".join(n.split())
    return difflib.SequenceMatcher(None, clean(a), clean(b)).ratio()


# ---------------------------------------------------------------- providers

class DemoCompanyProvider:
    """Synthetic MCA records derived from a hash of the PAN / CIN."""

    name = "demo"
    _ROCS = [("MH", "RoC-Mumbai"), ("DL", "RoC-Delhi"), ("KA", "RoC-Bangalore"), ("TN", "RoC-Chennai"),
             ("WB", "RoC-Kolkata"), ("GJ", "RoC-Ahmedabad"), ("HR", "RoC-Delhi"), ("TG", "RoC-Hyderabad")]
    _CITIES = {"MH": "Mumbai Maharashtra 400001", "DL": "New Delhi Delhi 110001", "KA": "Bengaluru Karnataka 560001",
               "TN": "Chennai Tamil Nadu 600001", "WB": "Kolkata West Bengal 700001",
               "GJ": "Ahmedabad Gujarat 380001", "HR": "Gurugram Haryana 122001", "TG": "Hyderabad Telangana 500001"}
    _STATUSES = (["Active"] * 14 + ["Strike Off"] * 3 + ["Under Process of Striking Off"]
                 + ["Dormant under section 455", "Amalgamated", "Under liquidation"])

    def lookup(self, identifier, id_type, name_hint=""):
        h = hashlib.sha256(identifier.encode()).digest()
        if id_type == "PAN" and h[0] % 10 == 0:
            return []
        records = [self._record(h, identifier, id_type, name_hint, self._STATUSES[h[1] % len(self._STATUSES)])]
        if id_type == "PAN" and h[2] % 12 == 0:  # same PAN shows up against two CINs
            twin = hashlib.sha256(h).digest()
            other = "Strike Off" if records[0]["status"] == "Active" else "Active"
            records.append(self._record(twin, identifier, id_type, name_hint, other))
        return records

    def _record(self, h, identifier, id_type, name_hint, status):
        code, roc = self._ROCS[h[3] % len(self._ROCS)]
        incorporated = date(1975, 1, 1) + timedelta(days=int.from_bytes(h[4:6], "big") % 17000)
        public = h[6] % 4 == 0
        listed = public and h[7] % 3 == 0
        reg_no = int.from_bytes(h[8:11], "big") % 900000 + 10000
        nic = int.from_bytes(h[11:13], "big") % 90000 + 10000
        cin = identifier if id_type == "CIN" else (
            f"{'L' if listed else 'U'}{nic}{code}{incorporated.year}{'PLC' if public else 'PTC'}{reg_no:06d}")
        name = (name_hint or f"DEMO COMPANY {cin[-6:]}").upper()
        name = re.sub(r"\bPVT\.?\b", "PRIVATE", name)
        name = re.sub(r"\bLTD\.?\b", "LIMITED", name)
        if not name.endswith("LIMITED"):
            name += " LIMITED" if public else " PRIVATE LIMITED"
        active = status == "Active"
        last_fy = 2025 if active else 2014 + h[13] % 8
        capital = (h[14] % 50 + 1) * 100000
        return {
            "cin": cin,
            "company_name": name,
            "status": status,
            "roc_code": roc,
            "registration_number": str(reg_no),
            "category": "Company limited by Shares",
            "subcategory": "Non-govt company",
            "class_of_company": "Public" if public else "Private",
            "paid_up_capital": str(capital),
            "number_of_members": "0",
            "date_of_incorporation": incorporated.strftime("%d/%m/%Y"),
            "registered_address": f"{h[15] % 200 + 1}, Industrial Area, {self._CITIES[code]}",
            "other_address": "",
            "email": f"accounts{h[16] % 90 + 10}@example.com",
            "listed": "Listed" if listed else "Unlisted",
            "suspended": "-",
            "active_compliance": "ACTIVE Compliant" if active else "ACTIVE non-compliant",
            "date_of_last_agm": f"{28 + h[17] % 3}/09/{last_fy}",
            "date_of_balance_sheet": f"31/03/{last_fy}",
        }


class LocalCompanyProvider:
    """Company master file (CSV / XLSX). Headers may be the report headers or field keys,
    plus a "PAN" column. Several rows with the same PAN produce a double status."""

    name = "local"

    def __init__(self, path):
        aliases = {label.lower(): key for key, label in COMPANY_COLUMNS}
        aliases.update({key: key for key in COMPANY_FIELDS})
        aliases.update({"company name": "company_name", "company status": "status", "pan": "pan",
                        "suspended at stock exchange": "suspended"})
        self._by_pan, self._by_cin = {}, {}
        for row in read_table(path):
            record = {}
            for header, value in row.items():
                key = aliases.get(str(header or "").strip().lower())
                if key and value not in (None, "") and key not in record:
                    record[key] = value.strftime("%d/%m/%Y") if hasattr(value, "strftime") else str(value).strip()
            if record.get("pan"):
                self._by_pan.setdefault(normalise(record["pan"]), []).append(record)
            if record.get("cin"):
                self._by_cin.setdefault(normalise(record["cin"]), []).append(record)

    def lookup(self, identifier, id_type, name_hint=""):
        index = self._by_pan if id_type == "PAN" else self._by_cin
        return [{f: r.get(f, "") for f in COMPANY_FIELDS} for r in index.get(identifier, [])]


class HttpCompanyProvider:
    """REST company-data API. MCA_API_URL_PAN / MCA_API_URL_CIN contain ``{id}``;
    MCA_API_RESULTS_PATH points at the record list (or single record) in the JSON;
    MCA_API_FIELD_MAP maps our field keys to dotted paths inside each record."""

    name = "http"

    def __init__(self, env=os.environ):
        self.url_pan = env.get("MCA_API_URL_PAN", "")
        self.url_cin = env.get("MCA_API_URL_CIN", "")
        if not (self.url_pan or self.url_cin):
            raise ProviderError("Set MCA_API_URL_PAN and/or MCA_API_URL_CIN for the http provider")
        self.method = env.get("MCA_API_METHOD", "GET").upper()
        self.timeout = float(env.get("MCA_API_TIMEOUT", "20"))
        self.headers = api_headers(env, "MCA")
        self.results_path = env.get("MCA_API_RESULTS_PATH", "")
        self.field_map = json.loads(env.get("MCA_API_FIELD_MAP", "{}")) or {f: f for f in COMPANY_FIELDS}

    def lookup(self, identifier, id_type, name_hint=""):
        template = self.url_pan if id_type == "PAN" else self.url_cin
        if not template:
            raise ProviderError(f"No API URL configured for {id_type} lookups")
        payload = fetch_json(template, identifier, self.method, self.headers, self.timeout)
        if payload is None:
            return []
        items = dig(payload, self.results_path) if self.results_path else payload
        if isinstance(items, dict):
            items = [items]
        records = []
        for item in items or []:
            rec = {f: dig(item, path) for f, path in self.field_map.items()}
            if rec.get("cin") or rec.get("company_name"):
                records.append({f: "" if rec.get(f) is None else str(rec[f]) for f in COMPANY_FIELDS})
        return records


def build_company_provider(env=os.environ):
    kind = env.get("MCA_PROVIDER", "demo").lower()
    if kind == "demo":
        return DemoCompanyProvider()
    if kind == "local":
        return LocalCompanyProvider(env.get("MCA_MASTER_FILE", "sample_data/company_master.csv"))
    if kind == "http":
        return HttpCompanyProvider(env)
    raise ProviderError(f"Unknown MCA_PROVIDER '{kind}' (use demo, local or http)")


# ------------------------------------------------------------------ service

class StruckOffService:
    def __init__(self, provider, cache=None):
        self.provider = provider
        self.cache = cache

    def _lookup(self, identifier, id_type, name_hint):
        key = f"mca:{self.provider.name}:{identifier}"
        cached = self.cache.get(key) if self.cache else None
        if cached is not None:
            return cached
        records = self.provider.lookup(identifier, id_type, name_hint)
        if self.cache:
            self.cache.put(key, records)
        return records

    def check(self, name="", gst="", pan="", cin="", sno=None):
        name = str(name or "").strip()
        gst, pan, cin = normalise(gst), normalise(pan), normalise(cin)
        result = {"sno": sno, "name": name, "gst": gst, "pan": pan, "cin_input": cin, "outcome": "",
                  "risk": "na", "status": "", "double_status": False, "records": [], "remarks": []}
        remarks = result["remarks"]

        if gst and not GSTIN_RE.match(gst):
            remarks.append("GSTIN format is invalid")
            gst_pan = ""
        else:
            gst_pan = gst[2:12] if gst else ""
        if pan and not PAN_RE.match(pan):
            remarks.append("PAN format is invalid")
            pan = ""
        if pan and gst_pan and pan != gst_pan:
            remarks.append(f"PAN does not match GSTIN (GSTIN contains {gst_pan})")
        if not pan and gst_pan:
            pan = gst_pan
            remarks.append("PAN taken from GSTIN")
        result["pan"] = pan or result["pan"]

        if cin and (CIN_RE.match(cin) or LLPIN_RE.match(cin)):
            identifier, id_type = cin, "CIN"
        elif pan:
            if pan[3] not in "CF":
                result.update(outcome=OUTCOME_NOT_COMPANY,
                              status=f"PAN holder: {pan_holder_type(pan)}")
                remarks.append("Only companies (PAN 4th letter C) and LLPs (F) are registered with MCA")
                return result
            identifier, id_type = pan, "PAN"
        else:
            result["outcome"] = OUTCOME_INVALID
            remarks.append("Provide a valid PAN, GSTIN or CIN")
            return result

        try:
            records = self._lookup(identifier, id_type, name)
        except ProviderError as exc:
            result["outcome"] = OUTCOME_ERROR
            remarks.append(str(exc))
            return result

        if not records:
            result["outcome"] = OUTCOME_NOT_FOUND
            result["status"] = OUTCOME_NOT_FOUND
            if pan[3:4] == "F":
                remarks.append("PAN is of a firm / LLP; a partnership firm is not registered with MCA")
            return result

        if name:  # best name match first, so it is used in the Final Sheet
            records = sorted(records, key=lambda r: -name_similarity(name, r.get("company_name")))
        primary = records[0]
        result["records"] = records
        result["double_status"] = len(records) > 1
        result["status"] = primary.get("status", "")
        result["outcome"], result["risk"] = classify_status(primary.get("status"))
        if result["double_status"]:
            others = sorted({r.get("status", "") for r in records[1:]})
            remarks.append(f"{len(records)} MCA records for this PAN (other status: {', '.join(others)})")
            worst = max((classify_status(r.get("status"))[1] for r in records),
                        key=["na", "ok", "medium", "high"].index)
            if worst != result["risk"] and result["risk"] == "ok":
                result["risk"] = "medium"
        if name and name_similarity(name, primary.get("company_name")) < 0.6:
            remarks.append("Name differs from MCA records - please review")
        if result["outcome"] == OUTCOME_ACTIVE and "non-compliant" in (primary.get("active_compliance") or "").lower():
            remarks.append("ACTIVE (INC-22A) non-compliant")
        return result

    def check_many(self, rows):
        return [self.check(sno=r.get("sno") or i, name=r.get("name"), gst=r.get("gst"),
                           pan=r.get("pan"), cin=r.get("cin"))
                for i, r in enumerate(rows, start=1)]


def summarise(results):
    order = [OUTCOME_ACTIVE, OUTCOME_STRUCK_OFF, OUTCOME_UNDER_PROCESS, OUTCOME_NOT_ACTIVE,
             OUTCOME_NOT_FOUND, OUTCOME_NOT_COMPANY, OUTCOME_INVALID, OUTCOME_ERROR]
    summary = {"Total suppliers": len(results)}
    for outcome in order:
        n = sum(1 for r in results if r["outcome"] == outcome)
        if n:
            summary[outcome] = n
    doubles = sum(1 for r in results if r["double_status"])
    if doubles:
        summary["Double status (multiple CINs)"] = doubles
    return summary


# -------------------------------------------------------------- excel input

def parse_supplier_file(filename, data):
    """Read a supplier list (the 'Suppliers' sheet format: S.No, Company Name, GST, PAN)."""
    lower = (filename or "").lower()
    if lower.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = next((wb[n] for n in wb.sheetnames if n.strip().lower() in ("suppliers", "supplier")), wb.active)
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    elif lower.endswith((".csv", ".txt")):
        import csv
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    else:
        raise ValueError("Upload an .xlsx or .csv file")

    header_idx, cols = None, {}
    for i, row in enumerate(rows[:10]):
        cells = [str(c or "").strip().lower() for c in row]
        found = {}
        for j, c in enumerate(cells):
            if "gst" in c and "gst" not in found:
                found["gst"] = j
            elif "pan" in c and "name" not in c and "pan" not in found:
                found["pan"] = j
            elif c.startswith("cin") or c == "llpin":
                found.setdefault("cin", j)
            elif "name" in c and "name" not in found:
                found["name"] = j
            elif c in ("s.no", "s.no.", "sno", "sr no", "sr. no.", "serial no", "s no"):
                found["sno"] = j
        if "pan" in found or "gst" in found or "cin" in found:
            header_idx, cols = i, found
            break
    if header_idx is None:
        raise ValueError("Could not find a PAN, GST or CIN column header")

    suppliers = []
    for row in rows[header_idx + 1:]:
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) and row[cols[k]] is not None else ""
        entry = {k: str(get(k)).strip() for k in ("name", "gst", "pan", "cin")}
        if not (entry["gst"] or entry["pan"] or entry["cin"]):
            continue
        sno = get("sno")
        entry["sno"] = sno if isinstance(sno, (int, float)) and not isinstance(sno, bool) else len(suppliers) + 1
        if isinstance(entry["sno"], float) and entry["sno"].is_integer():
            entry["sno"] = int(entry["sno"])
        suppliers.append(entry)
    return suppliers


# ------------------------------------------------------------- excel output

def _header(ws, labels, fill="1F4E79"):
    ws.append(labels)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _company_row(supplier, record):
    return [supplier["name"], supplier["pan"]] + [record.get(f, "") for f in COMPANY_FIELDS]


COMPANY_SHEET_HEADERS = ["Company Name", "PAN"] + [label for _, label in COMPANY_COLUMNS]
COMPANY_SHEET_WIDTHS = [45, 14, 24, 50, 28, 16, 20, 26, 22, 18, 18, 22, 20, 70, 50, 32, 20, 26, 22, 17, 20]
STATUS_COL = COMPANY_SHEET_HEADERS.index("Company Status(for efiling)") + 1


def build_struck_off_report(results):
    wb = Workbook()

    ws = wb.active
    ws.title = "Suppliers"
    _header(ws, ["S.No", "Company Name", "GST", "PAN", "MCA Status", "Result", "Remarks"])
    for r in results:
        ws.append([r.get("sno"), r["name"], r["gst"], r["pan"], r["status"], r["outcome"], "; ".join(r["remarks"])])
        fill = RISK_FILLS.get(r["risk"])
        if fill:
            ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=fill)
    _widths(ws, [6, 48, 20, 14, 30, 28, 60])

    final = wb.create_sheet("Supplier-Final Sheet")
    _header(final, COMPANY_SHEET_HEADERS)
    for r in results:
        record = r["records"][0] if r["records"] else {"status": r["status"] or r["outcome"]}
        final.append(_company_row(r, record))
        fill = RISK_FILLS.get(r["risk"])
        if fill:
            final.cell(final.max_row, STATUS_COL).fill = PatternFill("solid", fgColor=fill)
    _widths(final, COMPANY_SHEET_WIDTHS)

    double = wb.create_sheet("Supplier- Double Status")
    _header(double, COMPANY_SHEET_HEADERS)
    for r in results:
        if r["double_status"]:
            for record in r["records"]:
                double.append(_company_row(r, record))
                fill = RISK_FILLS.get(classify_status(record.get("status"))[1])
                if fill:
                    double.cell(double.max_row, STATUS_COL).fill = PatternFill("solid", fgColor=fill)
    _widths(double, COMPANY_SHEET_WIDTHS)

    summary = wb.create_sheet("Summary")
    summary.append(["Struck-Off Companies Report"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Generated on", datetime.now().strftime("%d/%m/%Y %H:%M")])
    summary.append([])
    summary.append(["Result", "Suppliers"])
    for cell in summary[4]:
        cell.font = Font(bold=True)
    for k, v in summarise(results).items():
        summary.append([k, v])
    _widths(summary, [34, 14])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_supplier_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Suppliers"
    _header(ws, ["S.No", "Company Name", "GST", "PAN"])
    ws.freeze_panes = "A2"
    ws.append([1, "Example Private Limited", "27AABCE1234F1Z5", "AABCE1234F"])
    ws.append([2, "Example Supplier (PAN only)", "", "AAACX5678K"])
    _widths(ws, [6, 48, 22, 16])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


SAMPLE_SUPPLIERS = [
    ("Sunrise Polymers Pvt Ltd", "27AAACS1234D1Z8", "AAACS1234D"),
    ("Navkar Logistics Private Limited", "", "AABCN4567P"),
    ("Apex Engineering Works Ltd", "24AADCA7788Q1ZM", ""),
    ("Bharat Packaging Pvt Ltd", "", "AAFCB2020K"),
    ("Om Sai Traders", "", "ABCPS1234K"),
    ("Pioneer Tech Solutions LLP", "", "AAKFP3344M"),
    ("Vertex Infra Private Limited", "07AAGCV9900H1Z2", "AAGCV9900H"),
    ("Shree Ganesh Industries Ltd", "", "AAACG5151R"),
    ("Metro Facility Services Pvt Ltd", "", "AACCM1019B"),
    ("Kaveri Agro Foods Private Limited", "29AAHCK8181J1Z0", "AAHCK8181J"),
    ("Delta Fabricators Pvt Ltd", "", "AADCD1019B"),
    ("Global Office Supplies", "", "AQZPK4321L"),
]
