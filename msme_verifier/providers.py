"""Data providers that resolve a PAN / Udyam number to MSME registration details.

The Udyam portal has no free public API, so live lookups go through a licensed
data provider (any REST API returning Udyam details). Three providers ship here:

* ``demo``  - deterministic synthetic data, for trying the tool without an API.
* ``local`` - looks identifiers up in a vendor master file (CSV / XLSX) you maintain.
* ``http``  - calls a configurable REST API (URL, auth header and field mapping via env).
"""
import csv
import hashlib
import json
import os
from datetime import date, timedelta

from .validators import PAN_HOLDER_TYPES, normalise

FIELDS = [
    "udyam_number",
    "pan",
    "enterprise_name",
    "enterprise_type",
    "major_activity",
    "organisation_type",
    "date_of_incorporation",
    "date_of_udyam_registration",
    "state",
    "district",
]


class ProviderError(Exception):
    """Raised when a provider cannot complete a lookup (network, auth, bad response)."""


def dig(data, path):
    """Follow a dotted path ("data.items.0.name") through nested dicts/lists."""
    for part in path.split("."):
        if isinstance(data, list) and part.isdigit():
            data = data[int(part)] if int(part) < len(data) else None
        elif isinstance(data, dict):
            data = data.get(part)
        else:
            return None
    return data


def api_headers(env, prefix):
    headers = {"Accept": "application/json"}
    if env.get(f"{prefix}_API_KEY"):
        headers[env.get(f"{prefix}_API_KEY_HEADER", "Authorization")] = env[f"{prefix}_API_KEY"]
    return headers


def fetch_json(url_template, identifier, method, headers, timeout):
    """Call a lookup API; returns parsed JSON, or None on HTTP 404."""
    import requests

    url = url_template.replace("{id}", identifier)
    try:
        if method == "POST":
            resp = requests.post(url, json={"id_number": identifier}, headers=headers, timeout=timeout)
        else:
            resp = requests.get(url, headers=headers, timeout=timeout)
    except requests.RequestException as exc:
        raise ProviderError(f"API request failed: {exc}") from exc
    if resp.status_code == 404:
        return None
    if resp.status_code >= 400:
        raise ProviderError(f"API returned HTTP {resp.status_code}")
    try:
        return resp.json()
    except ValueError as exc:
        raise ProviderError("API returned a non-JSON response") from exc


def read_table(path):
    """Rows of a CSV / XLSX file as dicts keyed by the header row."""
    if not os.path.exists(path):
        raise ProviderError(f"Master file not found: {path}")
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip() for h in rows[0]]
        return [dict(zip(headers, r)) for r in rows[1:]]
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


class BaseProvider:
    name = "base"

    def lookup(self, identifier, id_type):
        """Return a dict with keys from FIELDS, or None when not registered."""
        raise NotImplementedError


class DemoProvider(BaseProvider):
    """Synthetic, repeatable results derived from a hash of the identifier."""

    name = "demo"
    _STATES = [("MH", "Maharashtra", "Pune"), ("DL", "Delhi", "New Delhi"),
               ("KA", "Karnataka", "Bengaluru Urban"), ("GJ", "Gujarat", "Ahmedabad"),
               ("TN", "Tamil Nadu", "Chennai"), ("UP", "Uttar Pradesh", "Gautam Buddha Nagar")]
    _NAMES = ["Shree Ganesh", "Apex", "Sunrise", "Bharat", "Navkar", "Vertex", "Om Sai", "Pioneer"]
    _SUFFIX = ["Industries", "Engineering Works", "Traders", "Polymers", "Tech Solutions",
               "Fabricators", "Enterprises", "Packaging"]
    _ORG_BY_PAN = {"P": "Proprietary", "C": "Private Limited Company", "F": "Partnership",
                   "H": "Hindu Undivided Family", "T": "Trust", "A": "Society"}

    def lookup(self, identifier, id_type):
        h = hashlib.sha256(identifier.encode()).digest()
        if h[0] % 4 == 0:  # ~25% come back as not registered
            return None
        code, state, district = self._STATES[h[1] % len(self._STATES)]
        if id_type == "PAN":
            pan = identifier
            udyam = f"UDYAM-{code}-{h[2] % 40 + 1:02d}-{int.from_bytes(h[3:7], 'big') % 10**7:07d}"
        else:
            udyam = identifier
            letters = "".join(chr(65 + b % 26) for b in h[2:5])
            pan = f"{letters}{'PCF'[h[5] % 3]}{chr(65 + h[6] % 26)}{int.from_bytes(h[7:9], 'big') % 10000:04d}{chr(65 + h[9] % 26)}"
        incorporated = date(1995, 1, 1) + timedelta(days=int.from_bytes(h[10:12], "big") % 10000)
        registered = max(incorporated, date(2020, 7, 1)) + timedelta(days=h[12] * 4)
        return {
            "udyam_number": udyam,
            "pan": pan,
            "enterprise_name": f"{self._NAMES[h[13] % 8]} {self._SUFFIX[h[14] % 8]}".upper(),
            "enterprise_type": ["Micro", "Micro", "Small", "Medium"][h[15] % 4],
            "major_activity": ["Manufacturing", "Services", "Trading"][h[16] % 3],
            "organisation_type": self._ORG_BY_PAN.get(pan[3], "Others"),
            "date_of_incorporation": incorporated.isoformat(),
            "date_of_udyam_registration": registered.isoformat(),
            "state": state,
            "district": district,
        }


class LocalFileProvider(BaseProvider):
    """Looks up a vendor master file (CSV or XLSX) with one row per enterprise.

    Recognised column headers (case-insensitive) are the names in FIELDS, plus
    common variants such as "Udyam Registration Number" or "Name of Enterprise".
    """

    name = "local"
    _ALIASES = {
        "udyam_number": ["udyam_number", "udyam registration number", "udyam no", "udyam number", "urn"],
        "pan": ["pan", "pan number", "pan no"],
        "enterprise_name": ["enterprise_name", "name of enterprise", "enterprise name", "name"],
        "enterprise_type": ["enterprise_type", "type of enterprise", "category"],
        "major_activity": ["major_activity", "major activity", "major activity of enterprise", "activity"],
        "organisation_type": ["organisation_type", "type of organisation", "organization type"],
        "date_of_incorporation": ["date_of_incorporation", "date of incorporation", "date of commencement"],
        "date_of_udyam_registration": ["date_of_udyam_registration", "date of udyam registration", "registration date"],
        "state": ["state"],
        "district": ["district"],
    }

    def __init__(self, path):
        self.path = path
        self._by_pan, self._by_udyam = {}, {}
        for row in read_table(path):
            record = self._map_row(row)
            if record.get("pan"):
                self._by_pan[normalise(record["pan"])] = record
            if record.get("udyam_number"):
                self._by_udyam[normalise(record["udyam_number"])] = record

    def _map_row(self, row):
        lowered = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
        record = {}
        for field, aliases in self._ALIASES.items():
            for alias in aliases:
                value = lowered.get(alias)
                if value not in (None, ""):
                    record[field] = value.isoformat() if hasattr(value, "isoformat") else str(value).strip()
                    break
        return record

    def lookup(self, identifier, id_type):
        index = self._by_pan if id_type == "PAN" else self._by_udyam
        record = index.get(identifier)
        return {f: record.get(f, "") for f in FIELDS} if record else None


class HttpApiProvider(BaseProvider):
    """Generic REST provider, configured entirely through environment variables.

    MSME_API_URL_PAN / MSME_API_URL_UDYAM  URL templates containing ``{id}``
    MSME_API_METHOD                        GET (default) or POST (sends {"id_number": id})
    MSME_API_KEY, MSME_API_KEY_HEADER      auth header value / name (default "Authorization")
    MSME_API_FIELD_MAP                     JSON: our field -> dotted path in the response
    MSME_API_NOT_FOUND_PATH / _VALUE       dotted path + value meaning "not registered"
    """

    name = "http"

    def __init__(self, env=os.environ):
        self.url_pan = env.get("MSME_API_URL_PAN", "")
        self.url_udyam = env.get("MSME_API_URL_UDYAM", "")
        if not (self.url_pan or self.url_udyam):
            raise ProviderError("Set MSME_API_URL_PAN and/or MSME_API_URL_UDYAM for the http provider")
        self.method = env.get("MSME_API_METHOD", "GET").upper()
        self.timeout = float(env.get("MSME_API_TIMEOUT", "20"))
        self.headers = api_headers(env, "MSME")
        self.field_map = json.loads(env.get("MSME_API_FIELD_MAP", "{}")) or {f: f for f in FIELDS}
        self.not_found_path = env.get("MSME_API_NOT_FOUND_PATH", "")
        self.not_found_value = env.get("MSME_API_NOT_FOUND_VALUE", "")

    def lookup(self, identifier, id_type):
        template = self.url_pan if id_type == "PAN" else self.url_udyam
        if not template:
            raise ProviderError(f"No API URL configured for {id_type} lookups")
        payload = fetch_json(template, identifier, self.method, self.headers, self.timeout)
        if payload is None:
            return None
        if self.not_found_path and str(dig(payload, self.not_found_path)) == self.not_found_value:
            return None
        record = {f: dig(payload, path) for f, path in self.field_map.items()}
        if not (record.get("udyam_number") or record.get("enterprise_name")):
            return None
        return {f: "" if record.get(f) is None else str(record[f]) for f in FIELDS}


def build_provider(env=os.environ):
    kind = env.get("MSME_PROVIDER", "demo").lower()
    if kind == "demo":
        return DemoProvider()
    if kind == "local":
        return LocalFileProvider(env.get("MSME_MASTER_FILE", "sample_data/msme_master.csv"))
    if kind == "http":
        return HttpApiProvider(env)
    raise ProviderError(f"Unknown MSME_PROVIDER '{kind}' (use demo, local or http)")


__all__ = ["FIELDS", "PAN_HOLDER_TYPES", "ProviderError", "build_provider",
           "DemoProvider", "LocalFileProvider", "HttpApiProvider"]
