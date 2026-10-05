import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from struckoff_checker.app import create_app
from struckoff_checker.excel_io import extract_suppliers
from struckoff_checker.service import classify_status
from struckoff_checker.validators import classify_input, normalise_name, pan_from_gstin

ROOT = Path(__file__).resolve().parent.parent
MASTER = str(ROOT / "sample_data" / "struckoff_master.csv")
SUPPLIERS = ROOT / "sample_data" / "suppliers_to_check.csv"


@pytest.fixture
def client():
    env = {"STRUCKOFF_PROVIDER": "local", "STRUCKOFF_MASTER_FILE": MASTER, "STRUCKOFF_CACHE_DAYS": "0"}
    return create_app(env).test_client()


def test_validators():
    assert pan_from_gstin("27AAACA0242K1Z5") == "AAACA0242K"
    assert normalise_name("Aster Technologies Pvt. Ltd.") == "ASTER TECHNOLOGIES PRIVATE LIMITED"
    clean, errs = classify_input(gstin="27AAACA0242K1Z5")
    assert clean["pan"] == "AAACA0242K" and not errs
    assert classify_input(pan="ABC123")[1]
    assert classify_input(pan="AAACA0242K", gstin="27AABCS1111A1Z7")[1]  # PAN/GSTIN mismatch
    assert classify_input()[1]


def test_classify_status():
    assert classify_status("Strike Off") == "Struck-Off"
    assert classify_status("Under process of striking off") == "Under Strike-Off Process"
    assert classify_status("ACTIVE compliant") == "Active"
    assert classify_status("Amalgamated") == "Other Status"


def test_check_flags(client):
    def chk(**kw):
        return client.post("/api/check", json=kw).get_json()
    r = chk(name="Sunrise Polymers Pvt Ltd")
    assert r["flag"] == "Struck-Off" and r["matched_by"] == "NAME"
    assert chk(pan="AAACA0242K")["flag"] == "Active"
    assert chk(gstin="27AAACA0242K1Z5")["matched_by"] == "PAN"
    assert chk(cin="U51909TN2001PTC046000")["flag"] == "Other Status"
    d = chk(name="Pioneer Engineering Private Limited")
    assert d["flag"] == "Double Status" and len(d["records"]) == 2
    assert chk(name="Nobody Private Limited")["flag"] == "Not Found"
    assert chk(pan="ABC123")["flag"] == "Invalid Input"
    assert client.post("/api/check", json={}).status_code == 400


def test_bulk_and_export(client):
    resp = client.post("/api/bulk", data={"file": (io.BytesIO(SUPPLIERS.read_bytes()), "s.csv")},
                       content_type="multipart/form-data")
    body = resp.get_json()
    assert [r["flag"] for r in body["results"]] == [
        "Active", "Struck-Off", "Under Strike-Off Process", "Double Status", "Other Status", "Not Found", "Invalid Input"]
    exp = client.post("/api/export", json={"results": body["results"]})
    wb = load_workbook(io.BytesIO(exp.data))
    assert wb.sheetnames == ["Suppliers", "Supplier-Final Sheet", "Supplier- Double Status", "Sheet1", "Summary"]
    final = wb["Supplier-Final Sheet"]
    assert final.max_column == 23 and final["C1"].value == "CIN"
    assert wb["Supplier- Double Status"].max_row == 3  # header + 2 conflicting records
    assert wb["Sheet1"]["E2"].value in ("Match", "Similar", "Mismatch")
    assert client.post("/api/export", json={}).status_code == 400


def test_extract_suppliers_headers():
    data = b"s.no,Compamy Name,GST,PAN\n1,Acme Pvt Ltd,27AAACA0242K1Z5,AAACA0242K\n"
    assert extract_suppliers("x.csv", data) == [
        {"name": "Acme Pvt Ltd", "gstin": "27AAACA0242K1Z5", "pan": "AAACA0242K", "cin": ""}]


def test_template_and_index(client):
    assert client.get("/api/template").status_code == 200
    assert b"Struck-Off" in client.get("/").data


def test_demo_provider_runs():
    c = create_app({"STRUCKOFF_CACHE_DAYS": "0"}).test_client()
    r = c.post("/api/check", json={"name": "Some Company Pvt Ltd"}).get_json()
    assert r["flag"] in {"Active", "Struck-Off", "Under Strike-Off Process", "Other Status", "Double Status", "Not Found"}


# ---- data.gov.in loader + mca_db provider ----
import sqlite3

from struckoff_checker import ogd_loader


@pytest.fixture
def mca_db(tmp_path):
    path = str(tmp_path / "mca.sqlite3")
    assert ogd_loader.main([str(ROOT / "sample_data" / "ogd_sample_company_master.csv"), "-o", path]) == 0
    return path


def test_loader_maps_ogd_columns(mca_db):
    db = sqlite3.connect(mca_db)
    assert db.execute("SELECT COUNT(*) FROM companies").fetchone()[0] == 4  # row without CIN skipped
    row = db.execute("SELECT date_of_incorporation, paid_up_capital, registration_number, roc_code, last_agm "
                     "FROM companies WHERE cin='U72200MH2005PTC123456'").fetchone()
    assert row == ("12/03/2005", "1000000", "123456", "RoC-Mumbai", "2025")


def test_loader_is_idempotent_and_rejects_bad_file(mca_db, tmp_path):
    sample = str(ROOT / "sample_data" / "ogd_sample_company_master.csv")
    assert ogd_loader.main([sample, "-o", mca_db]) == 0  # reload: no duplicates
    assert sqlite3.connect(mca_db).execute("SELECT COUNT(*) FROM companies").fetchone()[0] == 4
    bad = tmp_path / "bad.csv"
    bad.write_text("a,b\n1,2\n")
    assert ogd_loader.main([str(bad), "-o", mca_db]) == 1


def test_mca_db_provider_end_to_end(mca_db):
    env = {"STRUCKOFF_PROVIDER": "mca_db", "STRUCKOFF_MCA_DB": mca_db, "STRUCKOFF_CACHE_DAYS": "0"}
    c = create_app(env).test_client()
    assert b"MCA open data loaded" in c.get("/").data

    def chk(**kw):
        return c.post("/api/check", json=kw).get_json()
    assert chk(name="Sunrise Polymers Pvt Ltd")["flag"] == "Struck-Off"
    assert chk(cin="U72200MH2005PTC123456")["flag"] == "Active"
    d = chk(name="Pioneer Engineering Pvt. Ltd.")
    assert d["flag"] == "Double Status" and len(d["records"]) == 2
    assert chk(name="Nobody Ltd")["flag"] == "Not Found"
    r = chk(pan="AAACA0242K")  # open data has no PAN -> reported, not silently ignored
    assert r["flag"] == "Not Found" and "PAN" in r["remarks"]


def test_mca_db_missing_file():
    from struckoff_checker.providers import ProviderError, build_provider
    with pytest.raises(ProviderError):
        build_provider({"STRUCKOFF_PROVIDER": "mca_db", "STRUCKOFF_MCA_DB": "/nonexistent.sqlite3"})
