import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from msme_verifier.app import create_app
from msme_verifier.providers import DemoProvider, HttpApiProvider, LocalFileProvider, ProviderError
from msme_verifier.service import VerificationService, payment_rule

ROOT = Path(__file__).resolve().parent.parent
MASTER = str(ROOT / "sample_data" / "msme_master.csv")


@pytest.fixture
def client():
    env = {"MSME_PROVIDER": "local", "MSME_MASTER_FILE": MASTER, "MSME_CACHE_DAYS": "0"}
    return create_app(env).test_client()


def test_verify_by_pan(client):
    r = client.post("/api/verify", json={"id": "aaapl1234c"}).get_json()
    assert r["status"] == "Registered"
    assert r["udyam_number"] == "UDYAM-MH-26-0012345"
    assert r["enterprise_name"] == "LAKSHMI PRECISION COMPONENTS"
    assert r["enterprise_type"] == "Micro"
    assert r["msme_payment_rule"] == "Yes"


def test_verify_by_udyam(client):
    r = client.post("/api/verify", json={"id": "UDYAM-GJ-01-0098765"}).get_json()
    assert r["status"] == "Registered"
    assert r["pan"] == "AADCS4321M"
    assert r["msme_payment_rule"] == "No"  # Medium enterprise


def test_not_registered_and_invalid(client):
    r = client.post("/api/verify", json={"id": "ABCPZ1111Z"}).get_json()
    assert r["status"] == "Not Registered"
    assert r["organisation_type"] == "Individual / Proprietorship"
    r = client.post("/api/verify", json={"id": "12345"}).get_json()
    assert r["status"] == "Invalid Format"
    assert client.post("/api/verify", json={"id": " "}).status_code == 400


def test_bulk_csv_upload_and_export(client):
    data = (ROOT / "sample_data" / "vendors_to_verify.csv").read_bytes()
    resp = client.post("/api/bulk", data={"file": (io.BytesIO(data), "vendors.csv")},
                       content_type="multipart/form-data")
    body = resp.get_json()
    assert resp.status_code == 200
    statuses = [r["status"] for r in body["results"]]
    assert statuses == ["Registered", "Registered", "Registered", "Registered", "Not Registered", "Invalid Format"]
    assert body["summary"]["Registered"] == 4
    trader = body["results"][3]
    assert trader["msme_payment_rule"] == "No (Trader)"

    export = client.post("/api/export", json={"results": body["results"]})
    assert export.status_code == 200
    wb = load_workbook(io.BytesIO(export.data))
    ws = wb["MSME Verification"]
    assert ws["C1"].value == "MSME Registration Status"
    assert ws.max_row == 7
    assert "Summary" in wb.sheetnames


def test_bulk_xlsx_without_header_and_paste(client):
    wb = Workbook()
    for v in ["AAAPL1234C", "UDYAM-DL-08-0045678", "AAAPL1234C"]:
        wb.active.append([v])
    buf = io.BytesIO()
    wb.save(buf)
    body = client.post("/api/bulk", data={"file": (io.BytesIO(buf.getvalue()), "x.xlsx")},
                       content_type="multipart/form-data").get_json()
    assert len(body["results"]) == 3
    assert all(r["status"] == "Registered" for r in body["results"])

    body = client.post("/api/bulk", json={"ids": "AAAPL1234C\nUDYAM-KA-03-0011122"}).get_json()
    assert len(body["results"]) == 2
    assert client.post("/api/bulk", json={"ids": []}).status_code == 400


def test_template_and_index(client):
    assert client.get("/api/template").status_code == 200
    html = client.get("/").get_data(as_text=True)
    assert "MSME Search by PAN" in html


def test_cache_serves_repeat_lookups(tmp_path):
    calls = []

    class CountingProvider(DemoProvider):
        def lookup(self, identifier, id_type):
            calls.append(identifier)
            return super().lookup(identifier, id_type)

    env = {"MSME_CACHE_DB": str(tmp_path / "c.sqlite3"), "MSME_CACHE_DAYS": "30"}
    client = create_app(env, provider=CountingProvider()).test_client()
    first = client.post("/api/verify", json={"id": "AAAPL1234C"}).get_json()
    second = client.post("/api/verify", json={"id": "AAAPL1234C"}).get_json()
    assert calls == ["AAAPL1234C"]
    assert second["remarks"] == "Served from cache"
    assert first["enterprise_name"] == second["enterprise_name"]


def test_demo_provider_is_deterministic():
    service = VerificationService(DemoProvider())
    a, b = service.verify("AAAPL1234C"), service.verify("AAAPL1234C")
    a.pop("checked_at"), b.pop("checked_at")
    assert a == b


def test_payment_rule():
    assert payment_rule("Micro", "Manufacturing") == "Yes"
    assert payment_rule("Small", "Services") == "Yes"
    assert payment_rule("Medium", "Services") == "No"
    assert payment_rule("Micro", "Trading") == "No (Trader)"
    assert payment_rule("", "") == "No"


def test_http_provider_mapping(monkeypatch):
    class FakeResp:
        status_code = 200

        @staticmethod
        def json():
            return {"data": {"urn": "UDYAM-MH-26-0012345", "name": "ACME", "category": "Small",
                             "activity": "Services", "pan": "AAAPL1234C"}}

    import requests
    monkeypatch.setattr(requests, "get", lambda url, **kw: FakeResp())
    provider = HttpApiProvider({
        "MSME_API_URL_PAN": "https://api.example.com/msme?pan={id}",
        "MSME_API_FIELD_MAP": '{"udyam_number": "data.urn", "enterprise_name": "data.name", '
                              '"enterprise_type": "data.category", "major_activity": "data.activity", "pan": "data.pan"}',
    })
    r = VerificationService(provider).verify("AAAPL1234C")
    assert r["status"] == "Registered"
    assert r["enterprise_name"] == "ACME"
    assert r["msme_payment_rule"] == "Yes"
    r = VerificationService(provider).verify("UDYAM-MH-26-0012345")
    assert r["status"] == "Error"  # no Udyam URL configured


def test_missing_master_file():
    with pytest.raises(ProviderError):
        LocalFileProvider("/nonexistent.csv")
