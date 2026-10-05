import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from msme_verifier import struck_off as so
from msme_verifier.app import create_app

ROOT = Path(__file__).resolve().parent.parent
MASTER = str(ROOT / "sample_data" / "company_master.csv")


@pytest.fixture
def client():
    env = {"MCA_PROVIDER": "local", "MCA_MASTER_FILE": MASTER, "MSME_CACHE_DAYS": "0"}
    return create_app(env).test_client()


def bulk_sample(client):
    data = (ROOT / "sample_data" / "suppliers_to_check.csv").read_bytes()
    resp = client.post("/api/struck-off/bulk", data={"file": (io.BytesIO(data), "s.csv")},
                       content_type="multipart/form-data")
    assert resp.status_code == 200
    return resp.get_json()


def test_bulk_outcomes(client):
    body = bulk_sample(client)
    by_sno = {r["sno"]: r for r in body["results"]}
    assert by_sno[1]["outcome"] == "Active"
    assert by_sno[2]["outcome"] == "Struck Off" and by_sno[2]["risk"] == "high"
    assert by_sno[3]["outcome"] == "Under Process of Striking Off"
    assert "PAN taken from GSTIN" in by_sno[3]["remarks"]
    assert by_sno[4]["double_status"] and by_sno[4]["outcome"] == "Active" and by_sno[4]["risk"] == "medium"
    assert by_sno[5]["outcome"] == "Active"  # LLP
    assert by_sno[6]["outcome"] == "Other - Not Active"
    assert by_sno[7]["outcome"] == "Not a Company / LLP"
    assert by_sno[8]["outcome"] == "Not Found on MCA"
    assert by_sno[9]["outcome"] == "Invalid Input"
    assert any("does not match GSTIN" in m for m in by_sno[10]["remarks"])
    assert body["summary"]["Total suppliers"] == 10


def test_report_format(client):
    results = bulk_sample(client)["results"]
    resp = client.post("/api/struck-off/report", json={"results": results})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Suppliers", "Supplier-Final Sheet", "Supplier- Double Status", "Summary"]
    assert [c.value for c in wb["Suppliers"][1]][:4] == ["S.No", "Company Name", "GST", "PAN"]
    final = wb["Supplier-Final Sheet"]
    headers = [c.value for c in final[1]]
    assert headers[:5] == ["Company Name", "PAN", "CIN", "Company Name As Per CIN", "Company Status(for efiling)"]
    assert headers[-1] == "Date of Balance Sheet" and len(headers) == 21
    assert final.max_row == 11  # one row per supplier
    assert final["E3"].value == "Strike Off"
    double = wb["Supplier- Double Status"]
    assert double.max_row == 3  # both CINs of supplier 4
    assert {double["C2"].value, double["C3"].value} == {"U21012KA2015PTC081122", "U21012KA2009PTC050011"}


def test_double_status_prefers_name_match(client):
    r = client.post("/api/struck-off/bulk", json={"suppliers": [
        {"name": "Bharat Packaging (South) Pvt Ltd", "pan": "AAFCB4444D"}]}).get_json()["results"][0]
    assert r["records"][0]["cin"] == "U21012KA2009PTC050011"
    assert r["outcome"] == "Struck Off"


def test_single_check_by_pan_gstin_cin(client):
    assert client.post("/api/struck-off/check", json={"id": "aaccn2222b"}).get_json()["outcome"] == "Struck Off"
    r = client.post("/api/struck-off/check", json={"id": "27AABCS1111A1Z5"}).get_json()
    assert r["outcome"] == "Active" and r["pan"] == "AABCS1111A"
    r = client.post("/api/struck-off/check", json={"id": "U45200TG2010PTC067890"}).get_json()
    assert r["status"] == "Dormant under section 455"
    assert client.post("/api/struck-off/check", json={"id": ""}).status_code == 400


def test_xlsx_upload_uses_suppliers_sheet():
    wb = Workbook()
    wb.active.title = "Other"
    ws = wb.create_sheet("Suppliers")
    ws.append(["s.no", "Compamy Name", "GST", "PAN"])
    ws.append([1, "Navkar Logistics", None, "AACCN2222B"])
    ws.append([2, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    rows = so.parse_supplier_file("x.xlsx", buf.getvalue())
    assert rows == [{"name": "Navkar Logistics", "gst": "", "pan": "AACCN2222B", "cin": "", "sno": 1}]
    with pytest.raises(ValueError):
        so.parse_supplier_file("x.pdf", b"")


def test_template_sample_and_page(client):
    assert client.get("/api/struck-off/template").status_code == 200
    sample = client.get("/api/struck-off/sample-report")
    wb = load_workbook(io.BytesIO(sample.data))
    assert wb["Supplier- Double Status"].max_row > 1
    assert "Struck-Off Companies" in client.get("/struck-off").get_data(as_text=True)


def test_classify_status():
    assert so.classify_status("Active") == ("Active", "ok")
    assert so.classify_status("STRIKE OFF") == ("Struck Off", "high")
    assert so.classify_status("Under Process of Striking Off") == ("Under Process of Striking Off", "high")
    assert so.classify_status("Amalgamated") == ("Other - Not Active", "medium")
    assert so.classify_status("Converted to LLP and Dissolved") == ("Other - Not Active", "high")


def test_http_provider(monkeypatch):
    class FakeResp:
        status_code = 200

        @staticmethod
        def json():
            return {"result": {"companies": [{"cin": "U1", "name": "ACME PRIVATE LIMITED", "status": "Strike Off"}]}}

    import requests
    monkeypatch.setattr(requests, "get", lambda url, **kw: FakeResp())
    provider = so.HttpCompanyProvider({
        "MCA_API_URL_PAN": "https://api.example.com/company?pan={id}",
        "MCA_API_RESULTS_PATH": "result.companies",
        "MCA_API_FIELD_MAP": '{"cin": "cin", "company_name": "name", "status": "status"}',
    })
    r = so.StruckOffService(provider).check(name="Acme Pvt Ltd", pan="AABCA1234B")
    assert r["outcome"] == "Struck Off"
    assert r["records"][0]["company_name"] == "ACME PRIVATE LIMITED"
